"""قراءة ملف الفواتير (Excel/CSV) واستخراج بيانات العملاء المُطبَّعة."""
import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .amounts import parse_amount_value
from .phone import PhoneResult, normalize_phone

CUSTOMER_INFO_RE = re.compile(r"^(?P<code>.+?)\s*-\s*(?P<name>.+?)\s*-\s*(?P<phone>[0-9٠-٩۰-۹\s\-\+]+)$")
_CODE_PREFIX_RE = re.compile(r"^[A-Za-z]+")

# حد أقصى لعدد الصفوف التي نفحصها بحثًا عن صف العناوين (صفوف عنوان الشركة/التقرير عادة قليلة جدًا)
MAX_HEADER_SCAN_ROWS = 30
# الحد الأدنى من الأعمدة المعروفة الواجب تواجدها بصف لاعتباره صف العناوين الفعلي
MIN_HEADER_MATCHES = 3


@dataclass
class InvoiceRecord:
    row_index: int
    customer_code: str
    customer_name: str
    phone_raw: str
    phone: PhoneResult
    invoice_date: "datetime | None"
    invoice_date_raw: str
    sales_order_id: str
    sales_type: str
    brand_raw: str
    brand_name_raw: str
    brand_code: "str | None"  # بادئة الكود المستخرجة (مثال: BQ) — للعرض فقط، غير مستخدم بالمطابقة
    store_id: str
    agent_created_by: str
    show_room_raw: str  # المصدر الفعلي لمطابقة البراند (انظر matcher.resolve_canonical_brand_from_show_room)
    amount: float  # من عمود "Total Amount" (أو "Amount" إن وُجد) — 0.0 إن كان العمود غير موجود بالملف


def find_column(df: pd.DataFrame, candidates: list[str]) -> "str | None":
    lower_map = {c.lower().strip(): c for c in df.columns}
    for cand in candidates:
        key = cand.lower().strip()
        if key in lower_map:
            return lower_map[key]
    return None


def parse_customer_info(text: str) -> "tuple[str, str, str] | None":
    """يفكك 'الكود - الاسم - الجوال' مع تحمّل أسماء تحتوي شرطات داخلية."""
    if not text or not isinstance(text, str):
        return None
    text = text.strip()
    parts = [p.strip() for p in text.split(" - ")]
    if len(parts) >= 3:
        code = parts[0]
        phone = parts[-1]
        name = " - ".join(parts[1:-1])
        return code, name, phone
    # fallback: ابحث عن رقم جوال في نهاية النص بأي تنسيق
    m = CUSTOMER_INFO_RE.match(text)
    if m:
        return m.group("code").strip(), m.group("name").strip(), m.group("phone").strip()
    return None


def extract_brand_code(customer_code: str, store_id: str) -> "str | None":
    """يستخرج بادئة الحروف من كود العميل، وإن فشلت يجرّب من Store ID."""
    for source in (customer_code, store_id):
        if not source:
            continue
        m = _CODE_PREFIX_RE.match(source.strip())
        if m:
            return m.group(0).upper()
    return None


def _parse_date(value) -> "datetime | None":
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if hasattr(value, "to_pydatetime"):
        try:
            return value.to_pydatetime()
        except Exception:
            pass
    text = str(value).strip()
    for fmt in ("%d %b %Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    try:
        return pd.to_datetime(text).to_pydatetime()
    except Exception:
        return None


def _flatten_column_candidates(invoice_columns_cfg: dict) -> set:
    candidates = set()
    for aliases in invoice_columns_cfg.values():
        for a in aliases:
            candidates.add(str(a).strip().lower())
    return candidates


def _score_row(row_values: list, candidates: set) -> int:
    return sum(1 for v in row_values if isinstance(v, str) and v.strip().lower() in candidates)


def _detect_header_row(peek_rows: list, candidates: set) -> "tuple[int, int]":
    """يرجع (أفضل رقم صف، عدد الأعمدة المطابقة به) بين الصفوف المفحوصة."""
    best_row, best_score = 0, -1
    for i, row in enumerate(peek_rows):
        score = _score_row(list(row), candidates)
        if score > best_score:
            best_score = score
            best_row = i
    return best_row, best_score


def _peek_csv_rows(file, max_rows: int) -> list:
    if hasattr(file, "seek"):
        file.seek(0)
    if hasattr(file, "read"):
        raw = file.read()
        text = raw.decode("utf-8-sig", errors="replace") if isinstance(raw, bytes) else raw
        if hasattr(file, "seek"):
            file.seek(0)
        source = io.StringIO(text)
        rows = []
        for i, row in enumerate(csv.reader(source)):
            if i >= max_rows:
                break
            rows.append(row)
        return rows
    with open(file, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
        rows = []
        for i, row in enumerate(csv.reader(f)):
            if i >= max_rows:
                break
            rows.append(row)
        return rows


def _peek_excel_rows(file, max_rows: int) -> list:
    if hasattr(file, "seek"):
        file.seek(0)
    wb = load_workbook(file, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append(["" if v is None else str(v) for v in row])
        return rows
    finally:
        wb.close()
        if hasattr(file, "seek"):
            file.seek(0)


def load_invoice_file(file, invoice_columns_cfg: dict, on_progress=None) -> pd.DataFrame:
    """
    يقبل مسار ملف (str/Path) أو كائن ملف بالذاكرة (مثل Streamlit UploadedFile الذي يحمل .name).

    صف العناوين (Header) الفعلي بالملف قد لا يكون السطر الأول — بعض التصديرات تضع صفوف
    عنوان الشركة/عنوان التقرير وصفوف فارغة قبله. لذلك نفحص أول MAX_HEADER_SCAN_ROWS صف
    ونختار الصف الذي يحوي أكبر عدد من أسماء الأعمدة المعروفة بـ invoice_columns، ثم نتجاهل
    كل ما قبله فعليًا عند القراءة النهائية.
    """
    log = on_progress or print
    name = getattr(file, "name", None) or str(file)
    suffix = Path(name).suffix.lower()
    is_excel = suffix in (".xlsx", ".xls")

    candidates = _flatten_column_candidates(invoice_columns_cfg)
    peek_rows = _peek_excel_rows(file, MAX_HEADER_SCAN_ROWS) if is_excel else _peek_csv_rows(file, MAX_HEADER_SCAN_ROWS)

    if not peek_rows:
        raise ValueError("الملف فارغ أو تعذّرت قراءته إطلاقًا — تحقق من أنه ليس تالفًا.")

    header_row_index, match_count = _detect_header_row(peek_rows, candidates)

    if match_count < MIN_HEADER_MATCHES:
        raise ValueError(
            f"تعذّر العثور تلقائيًا على صف العناوين (Header) ضمن أول {len(peek_rows)} صف بالملف. "
            f"أفضل صف تم فحصه هو رقم {header_row_index + 1} وبه {match_count} عمود معروف فقط "
            f"(الحد الأدنى المطلوب: {MIN_HEADER_MATCHES}). "
            f"تحقق من أن أسماء الأعمدة الفعلية بالملف مضافة كبديل ضمن invoice_columns بملف config.yaml، "
            f"أو أن صف العناوين موجود فعليًا ضمن أول {MAX_HEADER_SCAN_ROWS} صف بالملف (زد MAX_HEADER_SCAN_ROWS "
            f"بملف src/invoices.py إن لزم)."
        )

    log(f"[معلومة] تم اعتبار الصف رقم {header_row_index + 1} بالملف كصف العناوين (Header) "
        f"— تطابق {match_count} عمود معروف من config.yaml، وتم تجاهل كل الصفوف قبله.")

    if hasattr(file, "seek"):
        file.seek(0)

    if is_excel:
        return pd.read_excel(file, skiprows=header_row_index, header=0, dtype=str)
    return pd.read_csv(file, skiprows=header_row_index, header=0, dtype=str, keep_default_na=False)


def extract_invoice_records(df: pd.DataFrame, invoice_columns_cfg: dict) -> "tuple[list[InvoiceRecord], list[dict]]":
    """
    يرجع (records, parse_errors).
    parse_errors: صفوف فشل تفكيك عمود Customer Info فيها بالكامل (لا يوجد رقم جوال أصلًا لاستخراجه).
    """
    col = {}
    for key, candidates in invoice_columns_cfg.items():
        col[key] = find_column(df, candidates)

    missing_required = [k for k in ("customer_info", "date", "sales_order_id", "show_room") if col.get(k) is None]
    if missing_required:
        raise ValueError(
            f"لم يتم العثور على الأعمدة التالية بملف الفواتير: {missing_required}. "
            f"راجع invoice_columns بملف config.yaml وأضف اسم العمود الفعلي كبديل."
        )

    records: list[InvoiceRecord] = []
    parse_errors: list[dict] = []

    for idx, row in df.iterrows():
        customer_info_raw = row.get(col["customer_info"], "") if col["customer_info"] else ""
        parsed = parse_customer_info(customer_info_raw)
        if parsed is None:
            parse_errors.append({
                "row_index": idx,
                "customer_info_raw": customer_info_raw,
                "reason": "تعذّر تفكيك عمود Customer Info إلى كود/اسم/جوال",
            })
            continue

        code, name, phone_raw = parsed
        phone_result = normalize_phone(phone_raw)

        store_id = row.get(col["store_id"], "") if col.get("store_id") else ""
        brand_code = extract_brand_code(code, store_id)

        record = InvoiceRecord(
            row_index=idx,
            customer_code=code,
            customer_name=name,
            phone_raw=phone_raw,
            phone=phone_result,
            invoice_date=_parse_date(row.get(col["date"], "")),
            invoice_date_raw=str(row.get(col["date"], "")),
            sales_order_id=str(row.get(col["sales_order_id"], "")).strip(),
            sales_type=str(row.get(col["sales_type"], "")).strip() if col.get("sales_type") else "",
            brand_raw=str(row.get(col["brand"], "")).strip() if col.get("brand") else "",
            brand_name_raw=str(row.get(col["brand_name"], "")).strip() if col.get("brand_name") else "",
            brand_code=brand_code,
            store_id=str(store_id).strip(),
            agent_created_by=str(row.get(col["create_by"], "")).strip() if col.get("create_by") else "",
            show_room_raw=str(row.get(col["show_room"], "")).strip() if col.get("show_room") else "",
            amount=parse_amount_value(row.get(col["amount"], "")) if col.get("amount") else 0.0,
        )
        records.append(record)

    return records, parse_errors
